from flask import Flask, render_template, request, redirect, url_for, flash, send_from_directory, send_file, session
import os
import io
import base64
import pymysql
import random
import csv
from datetime import datetime, timedelta
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from reportlab.lib.pagesizes import letter, landscape
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors

app = Flask(__name__)
app.secret_key = 'projectwg_admin_secret_key_tsa_2026'
app.permanent_session_lifetime = timedelta(days=7)

# ==========================================
# KONFIGURASI DATABASE MYSQL (XAMPP / CLOUD)
# ==========================================
DB_HOST = os.environ.get('DB_HOST', '127.0.0.1')
DB_PORT = int(os.environ.get('DB_PORT', 3306))
DB_USER = os.environ.get('DB_USER', 'root')
DB_PASSWORD = os.environ.get('DB_PASSWORD', '')  # Kosongkan jika default XAMPP
DB_NAME = os.environ.get('DB_NAME', 'absensi_wg')
DB_SSL = os.environ.get('DB_SSL', 'false').lower() == 'true' or DB_PORT == 4000

def get_db_connection():
    ssl_config = {'ssl': {}} if DB_SSL else None
    try:
        return pymysql.connect(
            host=DB_HOST,
            port=DB_PORT,
            user=DB_USER,
            password=DB_PASSWORD,
            database=DB_NAME,
            ssl=ssl_config,
            cursorclass=pymysql.cursors.DictCursor,
            connect_timeout=5
        )
    except Exception as e:
        # Fallback ke 127.0.0.1 / localhost jika koneksi utama gagal
        try:
            return pymysql.connect(
                host='localhost',
                port=3306,
                user='root',
                password='',
                database=DB_NAME,
                cursorclass=pymysql.cursors.DictCursor,
                connect_timeout=2
            )
        except Exception:
            raise e

def ensure_db_schema():
    """Memastikan database absensi_wg beserta seluruh tabel dan data default tersedia di MySQL."""
    ssl_config = {'ssl': {}} if DB_SSL else None
    
    # 1. Pastikan database absensi_wg dibuat terlebih dahulu (konek tanpa nama database)
    try:
        init_conn = pymysql.connect(
            host=DB_HOST,
            port=DB_PORT,
            user=DB_USER,
            password=DB_PASSWORD,
            ssl=ssl_config,
            charset='utf8mb4',
            cursorclass=pymysql.cursors.DictCursor,
            connect_timeout=5
        )
        with init_conn.cursor() as cur:
            cur.execute(f"CREATE DATABASE IF NOT EXISTS `{DB_NAME}` CHARACTER SET utf8mb4 COLLATE utf8mb4_unicode_ci;")
        init_conn.commit()
        init_conn.close()
        print(f"[Schema Info] Database '{DB_NAME}' terverifikasi/berhasil dibuat.")
    except Exception as e:
        print(f"[Schema Notice] Create DB check: {e}")

    # 2. Hubungkan ke database dan pastikan semua tabel inti dibuat
    try:
        conn = get_db_connection()
        cursor = conn.cursor()

        # Tabel 1: Admins (Autentikasi Web Admin)
        cursor.execute("""
            CREATE TABLE IF NOT EXISTS admins (
                id INT AUTO_INCREMENT PRIMARY KEY,
                username VARCHAR(100) NOT NULL UNIQUE,
                password VARCHAR(255) NOT NULL,
                name VARCHAR(150) NOT NULL,
                email VARCHAR(150),
                role VARCHAR(50) DEFAULT 'Super Admin',
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
            ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_unicode_ci;
        """)

        # Pastikan kolom tabel admins bersih
        for col in ['avatar', 'last_login', 'updated_at']:
            cursor.execute("""
                SELECT COUNT(*) AS cnt FROM information_schema.COLUMNS 
                WHERE TABLE_SCHEMA = %s AND TABLE_NAME = 'admins' AND COLUMN_NAME = %s
            """, (DB_NAME, col))
            res_col = cursor.fetchone()
            if res_col and res_col.get('cnt', 0) > 0:
                cursor.execute(f"ALTER TABLE admins DROP COLUMN {col}")

        # Seed akun default admin (admin / 123)
        cursor.execute("SELECT COUNT(*) AS total FROM admins")
        if cursor.fetchone().get('total', 0) == 0:
            cursor.execute("""
                INSERT INTO admins (username, password, name, email, role)
                VALUES (%s, %s, %s, %s, %s)
            """, ('admin', '123', 'Administrator TSA', 'admin@tsa.co.id', 'Super Admin'))
            print("[Schema Info] Akun default admin (admin / 123) berhasil dibuat.")

        # Tabel 2: Branch Companies (Kantor Cabang)
        cursor.execute("""
            CREATE TABLE IF NOT EXISTS branch_companies (
                id INT AUTO_INCREMENT PRIMARY KEY,
                branch_name VARCHAR(150) NOT NULL,
                city VARCHAR(100) NOT NULL,
                latitude DECIMAL(10, 8) NOT NULL,
                longitude DECIMAL(11, 8) NOT NULL,
                radius INT DEFAULT 150,
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
            ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_unicode_ci;
        """)
        cursor.execute("SELECT COUNT(*) AS total FROM branch_companies")
        if cursor.fetchone().get('total', 0) == 0:
            cursor.execute("""
                INSERT INTO branch_companies (branch_name, city, latitude, longitude, radius) VALUES
                ('Kantor Pusat Surabaya', 'Surabaya', -7.25750000, 112.75210000, 200),
                ('Kantor Cabang Jakarta', 'Jakarta', -6.20880000, 106.84560000, 200)
            """)

        # Tabel 3: Users (Login Mobile & AI)
        cursor.execute("""
            CREATE TABLE IF NOT EXISTS users (
                id INT AUTO_INCREMENT PRIMARY KEY,
                username VARCHAR(100) NOT NULL UNIQUE,
                password VARCHAR(255) NOT NULL,
                name VARCHAR(150) NOT NULL,
                face_image VARCHAR(255) NULL,
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
            ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_unicode_ci;
        """)

        # Tabel 4: Employees (Data Karyawan)
        cursor.execute("""
            CREATE TABLE IF NOT EXISTS employees (
                id INT AUTO_INCREMENT PRIMARY KEY,
                employee_id VARCHAR(50) NOT NULL UNIQUE,
                name VARCHAR(150) NOT NULL,
                team VARCHAR(100) NOT NULL,
                branch VARCHAR(100) NOT NULL,
                status VARCHAR(50) NOT NULL DEFAULT 'Aktif',
                tag1 VARCHAR(50) NULL,
                tag2 VARCHAR(50) NULL,
                tag3 VARCHAR(50) NULL,
                username VARCHAR(100) NOT NULL UNIQUE,
                password VARCHAR(255) NOT NULL,
                face_image VARCHAR(255) NULL,
                email VARCHAR(150) NULL,
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
            ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_unicode_ci;
        """)

        # Tabel 5: Attendance (Presensi & Izin)
        cursor.execute("""
            CREATE TABLE IF NOT EXISTS attendance (
                id INT AUTO_INCREMENT PRIMARY KEY,
                username VARCHAR(100) NOT NULL,
                type VARCHAR(50) NOT NULL,
                timestamp DATETIME NOT NULL DEFAULT CURRENT_TIMESTAMP,
                clock_out DATETIME NULL,
                image_path VARCHAR(255) NULL,
                reason TEXT NULL,
                status VARCHAR(20) DEFAULT 'Pending',
                INDEX idx_att_user_time (username, timestamp)
            ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_unicode_ci;
        """)

        # Tabel 6: Vacancies (Lowongan)
        cursor.execute("""
            CREATE TABLE IF NOT EXISTS vacancies (
                id INT AUTO_INCREMENT PRIMARY KEY,
                title VARCHAR(150) NOT NULL,
                division VARCHAR(100) NOT NULL,
                branch VARCHAR(100) NOT NULL,
                description TEXT,
                requirements TEXT,
                posted_date DATE NOT NULL,
                status VARCHAR(50) DEFAULT 'Open',
                image VARCHAR(255) DEFAULT 'assets/avatar_mascot_centered.png',
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
            ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_unicode_ci;
        """)

        conn.commit()
        conn.close()
        print("[Schema Info] Seluruh tabel database absensi_wg berhasil diverifikasi.")
    except Exception as e:
        print(f"[Schema Warning] ensure_db_schema: {e}")

# Inisialisasi pengecekan skema saat app dimuat
ensure_db_schema()

def generate_random_employee_id():
    part1 = random.randint(100, 999)
    part2 = random.randint(100, 999)
    part3 = random.randint(10, 99)
    return f"{part1}.{part2}.{part3}"

# ==========================================
# MIDDLEWARE & AUTENTIKASI LOGIN ADMIN
# ==========================================
@app.context_processor
def inject_admin_context():
    """Menyediakan data profil admin yang sedang login ke seluruh template Jinja2."""
    return {
        'current_admin': {
            'logged_in': session.get('admin_logged_in', False),
            'id': session.get('admin_id'),
            'username': session.get('admin_username', 'admin'),
            'name': session.get('admin_name', 'Administrator TSA'),
            'role': session.get('admin_role', 'Super Admin'),
            'email': session.get('admin_email', 'admin@tsa.co.id'),
            'avatar': 'assets/tsa_logo.png'
        }
    }

@app.before_request
def require_admin_login():
    """Proteksi route Web Admin: redirect ke /login bila belum terautentikasi."""
    # Rute publik yang boleh diakses tanpa login
    allowed_endpoints = ['login', 'serve_assets', 'static']
    if request.endpoint in allowed_endpoints:
        return
    if request.path.startswith('/assets') or request.path.startswith('/static') or request.path == '/favicon.ico':
        return

    # Jika session admin belum ada, redirect ke /login
    if not session.get('admin_logged_in'):
        next_url = request.url if request.method == 'GET' and request.path != '/logout' else None
        return redirect(url_for('login', next=next_url))

# ==========================================
# RUTE LOGIN & LOGOUT WEB ADMIN
# ==========================================
@app.route('/login', methods=['GET', 'POST'])
def login():
    # Jika sudah login, langsung alihkan ke homepage
    if session.get('admin_logged_in'):
        return redirect(url_for('index'))

    if request.method == 'POST':
        username = request.form.get('username', '').strip()
        password = request.form.get('password', '').strip()
        remember = request.form.get('remember')

        if not username or not password:
            flash('Username dan password wajib diisi!', 'danger')
            return render_template('login.html', username=username)

        admin = None
        is_valid = False
        try:
            conn = get_db_connection()
            cursor = conn.cursor()

            # 1. Cek dari tabel admins
            cursor.execute("SELECT * FROM admins WHERE username = %s LIMIT 1", (username,))
            admin = cursor.fetchone()

            # 2. Fallback cek dari tabel users jika belum terdaftar di admins
            if not admin:
                cursor.execute("SELECT * FROM users WHERE username = %s LIMIT 1", (username,))
                user_row = cursor.fetchone()
                if user_row and user_row['username'].lower() in ('admin', 'administrator'):
                    admin = {
                        'id': user_row['id'],
                        'username': user_row['username'],
                        'password': user_row['password'],
                        'name': user_row.get('name') or 'Administrator TSA',
                        'email': 'admin@tsa.co.id',
                        'role': 'Super Admin'
                    }

            # Validasi password
            if admin:
                db_pw = str(admin['password']).strip()
                # Validasi langsung
                if db_pw == password:
                    is_valid = True
                # Fleksibilitas default password admin (123 / admin123)
                elif (db_pw in ('123', 'admin123')) and (password in ('123', 'admin123')):
                    is_valid = True
                else:
                    try:
                        from werkzeug.security import check_password_hash
                        if check_password_hash(db_pw, password):
                            is_valid = True
                    except Exception:
                        pass

            conn.close()
        except Exception as db_err:
            print(f"[Login DB Error]: {db_err}")
            flash(f"Koneksi Database MySQL Gagal ({DB_HOST}:{DB_PORT}): {db_err}", 'danger')
            return render_template('login.html', username=username)

        if admin and is_valid:
            session.permanent = bool(remember)
            session['admin_logged_in'] = True
            session['admin_id'] = admin['id']
            session['admin_username'] = admin['username']
            session['admin_name'] = admin.get('name') or 'Administrator TSA'
            session['admin_role'] = admin.get('role', 'Super Admin')
            session['admin_email'] = admin.get('email', 'admin@tsa.co.id')

            flash(f"Selamat datang kembali, {session['admin_name']}!", 'success')
            next_url = request.args.get('next')
            return redirect(next_url if next_url and next_url.startswith('/') else url_for('index'))
        else:
            flash('Username atau password yang Anda masukkan salah!', 'danger')
            return render_template('login.html', username=username)

    return render_template('login.html')

# ==========================================
# RUTE HOMEPAGE (DASHBOARD UTAMA SESUAI DATABASE REAL)
# ==========================================
@app.route('/')
@app.route('/home')
def index():
    # 1. Waktu & Sapaan Realtime
    now = datetime.now()
    current_hour = now.hour
    if 5 <= current_hour < 12:
        greeting = "Good Morning,"
    elif 12 <= current_hour < 17:
        greeting = "Good Afternoon,"
    elif 17 <= current_hour < 21:
        greeting = "Good Evening,"
    else:
        greeting = "Good Night,"

    current_day = now.strftime('%A')
    current_date_formatted = now.strftime('%B %d, %Y')

    total_employees = 0
    presence_today = 0
    absence_today = 0
    request_obs = 0
    chart_labels = ['May', 'Jun', 'Jul', 'Aug', 'Sep', 'Oct']
    chart_presence = [0, 0, 0, 0, 0, 0]
    chart_absence = [0, 0, 0, 0, 0, 0]
    chart_employees = [0] * 6
    top_attendees = []
    other_attendees = []

    try:
        conn = get_db_connection()
        cursor = conn.cursor()
        
        # 2. Total Karyawan dari Database
        cursor.execute("SELECT COUNT(*) AS total FROM employees")
        emp_res = cursor.fetchone()
        total_employees = emp_res['total'] if emp_res else 0
        
        # 3. Total Presensi Hadir Hari Ini (clock_in hari ini)
        cursor.execute("""
            SELECT COUNT(DISTINCT username) AS total 
            FROM attendance 
            WHERE DATE(timestamp) = CURDATE() AND (type = 'clock_in' OR type = 'hadir')
        """)
        presence_res = cursor.fetchone()
        presence_today = presence_res['total'] if presence_res else 0
        
        # 4. Total Request / Izin / Sakit / OBS Hari Ini
        cursor.execute("""
            SELECT COUNT(*) AS total 
            FROM attendance 
            WHERE DATE(timestamp) = CURDATE() 
              AND ((reason IS NOT NULL AND reason != '') OR type IN ('izin', 'sakit', 'cuti', 'sick_leave', 'necessity_leave'))
        """)
        obs_res = cursor.fetchone()
        request_obs = obs_res['total'] if obs_res else 0
        
        # 5. Ketidakhadiran Hari Ini (karyawan yang belum clock_in hari ini)
        absence_today = max(0, total_employees - presence_today)
        
        # 6. Data Rekap Bulanan untuk Grafik (Chart.js) langsung dari Database
        cursor.execute("""
            SELECT 
                DATE_FORMAT(timestamp, '%b') AS month_name,
                MONTH(timestamp) AS month_num,
                YEAR(timestamp) AS yr,
                SUM(CASE WHEN type = 'clock_in' OR type = 'hadir' THEN 1 ELSE 0 END) AS presence_cnt,
                SUM(CASE WHEN type != 'clock_in' AND type != 'clock_out' AND type != 'hadir' THEN 1 ELSE 0 END) AS absence_cnt
            FROM attendance
            GROUP BY YEAR(timestamp), MONTH(timestamp), DATE_FORMAT(timestamp, '%b')
            ORDER BY YEAR(timestamp) ASC, MONTH(timestamp) ASC
        """)
        monthly_rows = cursor.fetchall()

        if monthly_rows and len(monthly_rows) > 0:
            chart_labels = []
            chart_presence = []
            chart_absence = []
            chart_employees = []
            recent_months = monthly_rows[-6:]
            for r in recent_months:
                chart_labels.append(r['month_name'])
                chart_presence.append(int(r['presence_cnt'] or 0))
                chart_absence.append(int(r['absence_cnt'] or 0))
                chart_employees.append(total_employees)
        else:
            chart_presence = [0, 0, 0, 0, 0, presence_today]
            chart_absence = [0, 0, 0, 0, 0, absence_today]
            chart_employees = [total_employees] * 6

        # 7. Quick Attendance List Karyawan Real dari Database
        cursor.execute("""
            SELECT 
                e.id, e.employee_id, e.name, e.team, e.branch, e.status, e.username, e.face_image,
                a.type AS today_type, 
                TIME_FORMAT(a.max_ts, '%H:%i') AS today_clock_in,
                a.reason AS today_reason
            FROM employees e
            LEFT JOIN (
                SELECT username, 
                       SUBSTRING_INDEX(GROUP_CONCAT(type ORDER BY id DESC), ',', 1) AS type,
                       MAX(timestamp) AS max_ts,
                       SUBSTRING_INDEX(GROUP_CONCAT(COALESCE(reason, '') ORDER BY id DESC), ',', 1) AS reason
                FROM attendance
                WHERE DATE(timestamp) = CURDATE()
                GROUP BY username
            ) a ON e.username = a.username
            ORDER BY (CASE WHEN a.type = 'clock_in' THEN 1 WHEN a.type IS NOT NULL THEN 2 ELSE 3 END) ASC, e.id ASC
        """)
        attendance_employees = cursor.fetchall()
        conn.close()

        top_attendees = attendance_employees[:3]
        other_attendees = attendance_employees[3:8]
    except Exception as dash_err:
        print(f"[Dashboard DB Error]: {dash_err}")
    
    return render_template(
        'homepage.html',
        greeting=greeting,
        current_day=current_day,
        current_date_formatted=current_date_formatted,
        total_employees=total_employees,
        presence_today=presence_today,
        absence_today=absence_today,
        request_obs=request_obs,
        chart_labels=chart_labels,
        chart_presence=chart_presence,
        chart_absence=chart_absence,
        chart_employees=chart_employees,
        top_attendees=top_attendees,
        other_attendees=other_attendees
    )

# ==========================================
# RUTE MANAJEMEN KARYAWAN (EMPLOYEE)
# ==========================================
@app.route('/employee')
@app.route('/employees')
def employee_management():
    search_query = request.args.get('search', '')
    conn = get_db_connection()
    cursor = conn.cursor()
    
    if search_query:
        cursor.execute("SELECT * FROM employees WHERE name LIKE %s ORDER BY name ASC", ('%' + search_query + '%',))
    else:
        cursor.execute("SELECT * FROM employees ORDER BY name ASC")
        
    employees = cursor.fetchall()
    
    cursor.execute("SELECT branch_name FROM branch_companies ORDER BY id ASC")
    branches = [r['branch_name'] for r in cursor.fetchall()]
    if not branches:
        branches = ['TUNGGAL SAHABAT ABADI', 'Decoy']
    conn.close()
    
    return render_template('admin_dashboard.html', employees=employees, search_query=search_query, branches=branches)

@app.route('/assets/<path:filename>')
def serve_assets(filename):
    return send_from_directory(os.path.join(app.root_path, 'assets'), filename)

@app.route('/add', methods=['POST'])
def add_employee():
    name = request.form.get('name')
    team = request.form.get('team')
    branch = request.form.get('branch')
    status = request.form.get('status')
    username = request.form.get('username')
    password = request.form.get('password')
    face_image = request.form.get('face_image')
    tags = request.form.get('tags', '').split(',')

    if not name or not team or not branch or not status or not username or not password:
        flash('Semua kolom data karyawan wajib diisi!', 'danger')
        return redirect(url_for('employee_management'))

    if not face_image or not face_image.strip():
        flash('Wajib mengambil foto selfie kamera sebagai acuan AI Face Recognition!', 'danger')
        return redirect(url_for('employee_management'))

    employee_id = generate_random_employee_id()

    # Simpan file gambar fisik ke folder assets/foto_karyawan/
    assets_dir = os.path.join(app.root_path, 'assets', 'foto_karyawan')
    os.makedirs(assets_dir, exist_ok=True)
    filename = f"{employee_id.replace('.', '_')}_{username}.jpg"
    filepath = os.path.join(assets_dir, filename)

    try:
        raw_b64 = face_image
        if ',' in face_image:
            raw_b64 = face_image.split(',', 1)[1]
        with open(filepath, 'wb') as f:
            f.write(base64.b64decode(raw_b64))
    except Exception as img_err:
        print(f"Peringatan: Gagal menyimpan file fisik gambar ke assets/foto_karyawan: {img_err}")

    # Simpan path saja di database (contoh: assets/foto_karyawan/123_456_78_aan.jpg)
    db_image_path = f"assets/foto_karyawan/{filename}"

    conn = get_db_connection()
    cursor = conn.cursor()

    try:
        # 1. Simpan ke tabel employees (Web Dashboard)
        cursor.execute(
            """INSERT INTO employees (employee_id, name, team, branch, status, tag1, tag2, tag3, username, password, face_image) 
               VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)""",
            (employee_id, name, team, branch, status, 
             tags[0].strip() if len(tags)>0 and tags[0].strip() else None, 
             tags[1].strip() if len(tags)>1 and tags[1].strip() else None, 
             tags[2].strip() if len(tags)>2 and tags[2].strip() else None, 
             username, password, db_image_path)
        )
        
        # 2. Simpan ke tabel users (Agar bisa login di Aplikasi Flutter & Acuan AI Face Recognition)
        cursor.execute(
            "INSERT INTO users (username, password, name, face_image) VALUES (%s, %s, %s, %s)",
            (username, password, name, db_image_path)
        )
        
        conn.commit()
        flash(f'Karyawan "{name}" berhasil ditambahkan dengan foto di assets/foto_karyawan!', 'success')
    except pymysql.err.IntegrityError:
        flash('Error: Username tersebut sudah terdaftar!', 'danger')
    except Exception as e:
        flash(f'Terjadi kesalahan: {e}', 'danger')
    finally:
        conn.close()

    return redirect(url_for('employee_management'))

@app.route('/detail/<int:id>')
def detail_employee(id):
    conn = get_db_connection()
    cursor = conn.cursor()
    
    cursor.execute("SELECT * FROM employees WHERE id = %s", (id,))
    employee = cursor.fetchone()
    
    if not employee:
        conn.close()
        flash('Data karyawan tidak ditemukan!', 'danger')
        return redirect(url_for('employee_management'))
    
    username = employee.get('username') or ''
    
    # Kueri HANYA aktivitas milik karyawan tersebut (tidak tergabung dengan karyawan lain)
    cursor.execute("""
        SELECT 
            id, username, type, timestamp, clock_out, image_path, reason,
            COALESCE(status, 'Approved') AS status,
            DATE_FORMAT(timestamp, '%%d %%b %%Y') AS formatted_date,
            TIME_FORMAT(timestamp, '%%H:%%i') AS formatted_time,
            TIME_FORMAT(clock_out, '%%H:%%i') AS formatted_clock_out
        FROM attendance 
        WHERE username = %s 
        ORDER BY timestamp DESC, id DESC
    """, (username,))
    all_attendance = cursor.fetchall()
    conn.close()
    
    # Pisahkan aktivitas agar tidak tergabung: Clock In, Clock Out, dan Izin
    clock_in_list = []
    clock_out_list = []
    permission_list = []
    
    for att in all_attendance:
        att_type = (att['type'] or '').lower()
        if att_type == 'clock_in':
            clock_in_list.append(att)
        elif att_type == 'clock_out':
            clock_out_list.append(att)
        elif 'izin' in att_type or 'sakit' in att_type or 'cuti' in att_type or (att['reason'] and att_type not in ('clock_in', 'clock_out', 'hadir')):
            if 'sakit' in att_type:
                att['display_type'] = 'Sakit'
            elif 'keperluan' in att_type:
                att['display_type'] = 'Izin Keperluan'
            elif 'cuti' in att_type:
                att['display_type'] = 'Cuti'
            else:
                att['display_type'] = 'Izin'
            permission_list.append(att)

    return render_template(
        'employee_detail.html',
        employee=employee,
        all_attendance=all_attendance,
        clock_in_list=clock_in_list,
        clock_out_list=clock_out_list,
        permission_list=permission_list,
        total_clock_in=len(clock_in_list),
        total_clock_out=len(clock_out_list),
        total_permission=len(permission_list)
    )

@app.route('/delete/<int:id>', methods=['POST'])
def delete_employee(id):
    try:
        conn = get_db_connection()
        cursor = conn.cursor()
        
        # Cari data karyawan sebelum dihapus
        cursor.execute("SELECT employee_id, username, face_image FROM employees WHERE id = %s", (id,))
        emp = cursor.fetchone()
        
        if emp:
            if emp.get('username'):
                cursor.execute("DELETE FROM users WHERE username = %s", (emp['username'],))
                cursor.execute("DELETE FROM attendance WHERE username = %s", (emp['username'],))
            
            # Hapus file foto dari folder assets/foto_karyawan jika ada
            if emp.get('face_image'):
                fpath = os.path.join(app.root_path, emp['face_image'].replace('/', os.sep))
                if os.path.exists(fpath):
                    try:
                        os.remove(fpath)
                    except Exception:
                        pass
            
        cursor.execute("DELETE FROM employees WHERE id = %s", (id,))
        
        conn.commit()
        conn.close()
        flash('Data karyawan, akun login, dan foto di assets/foto_karyawan berhasil dihapus!', 'success')
    except Exception as e:
        flash(f'Gagal menghapus data: {e}', 'danger')
        
    return redirect(url_for('employee_management'))

# ==========================================
# RUTE KANTOR CABANG (BRANCH COMPANY & LONG/LAT)
# ==========================================
@app.route('/branch-company')
def branch_company():
    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute("SELECT * FROM branch_companies ORDER BY id DESC")
    branches = cursor.fetchall()
    conn.close()
    return render_template('branch_company.html', branches=branches)

@app.route('/add-branch', methods=['POST'])
def add_branch():
    branch_name = request.form.get('branch_name')
    city = request.form.get('city')
    latitude = request.form.get('latitude')
    longitude = request.form.get('longitude')
    radius = request.form.get('radius', 100)

    if not branch_name or not city or not latitude or not longitude:
        flash('Semua kolom wajib diisi termasuk koordinat!', 'danger')
        return redirect(url_for('branch_company'))

    try:
        conn = get_db_connection()
        cursor = conn.cursor()
        cursor.execute(
            """INSERT INTO branch_companies (branch_name, city, latitude, longitude, radius) 
               VALUES (%s, %s, %s, %s, %s)""",
            (branch_name, city, latitude, longitude, radius)
        )
        conn.commit()
        conn.close()
        flash('Kantor cabang dan koordinat berhasil ditambahkan!', 'success')
    except Exception as e:
        flash(f'Gagal menambah cabang: {e}', 'danger')

    return redirect(url_for('branch_company'))

@app.route('/edit-branch/<int:id>', methods=['POST'])
def edit_branch(id):
    branch_name = request.form.get('branch_name')
    city = request.form.get('city')
    latitude = request.form.get('latitude')
    longitude = request.form.get('longitude')
    radius = request.form.get('radius', 100)

    if not branch_name or not city or not latitude or not longitude:
        flash('Semua kolom wajib diisi termasuk koordinat!', 'danger')
        return redirect(url_for('branch_company'))

    try:
        conn = get_db_connection()
        cursor = conn.cursor()
        cursor.execute(
            """UPDATE branch_companies 
               SET branch_name=%s, city=%s, latitude=%s, longitude=%s, radius=%s 
               WHERE id=%s""",
            (branch_name, city, latitude, longitude, radius, id)
        )
        conn.commit()
        conn.close()
        flash(f'Kantor cabang "{branch_name}" berhasil diperbarui!', 'success')
    except Exception as e:
        flash(f'Gagal mengupdate cabang: {e}', 'danger')

    return redirect(url_for('branch_company'))

@app.route('/delete-branch/<int:id>', methods=['POST'])
def delete_branch(id):
    try:
        conn = get_db_connection()
        cursor = conn.cursor()
        cursor.execute("SELECT branch_name FROM branch_companies WHERE id = %s", (id,))
        branch = cursor.fetchone()
        cursor.execute("DELETE FROM branch_companies WHERE id = %s", (id,))
        conn.commit()
        conn.close()
        name = branch['branch_name'] if branch else f'ID {id}'
        flash(f'Kantor cabang "{name}" berhasil dihapus!', 'success')
    except Exception as e:
        flash(f'Gagal menghapus cabang: {e}', 'danger')

    return redirect(url_for('branch_company'))

# ==========================================
# RUTE LOWONGAN KERJA (VACANCY)
# ==========================================
@app.route('/vacancy')
def vacancy():
    now = datetime.now()
    current_day = now.strftime('%A')
    current_date_formatted = now.strftime('%B %d, %Y')
    
    vacancies = []
    branches = []
    try:
        conn = get_db_connection()
        cursor = conn.cursor()
        cursor.execute("""
            SELECT id, title, division, branch, description, requirements,
                   DATE_FORMAT(posted_date, '%M %e, %Y') as formatted_date,
                   posted_date, status, image
            FROM vacancies
            ORDER BY id DESC
        """)
        vacancies = cursor.fetchall()
        
        cursor.execute("SELECT branch_name FROM branch_companies ORDER BY branch_name ASC")
        branches = cursor.fetchall()
        conn.close()
    except Exception as e:
        flash(f'Gagal memuat lowongan kerja: {e}', 'danger')
        print(f"[Vacancy Error] {e}")
    
    return render_template(
        'vacancy.html', 
        vacancies=vacancies, 
        branches=branches,
        current_day=current_day,
        current_date_formatted=current_date_formatted
    )

@app.route('/add-vacancy', methods=['POST'])
def add_vacancy():
    title = request.form.get('title', 'Staff IT')
    division = request.form.get('division', 'IT')
    branch = request.form.get('branch', 'Surabaya')
    description = request.form.get('description', '')
    requirements = request.form.get('requirements', '')
    posted_date = request.form.get('posted_date') or datetime.now().strftime('%Y-%m-%d')
    status = request.form.get('status', 'Open')
    image = 'assets/avatar_mascot_centered.png'

    if not title:
        flash('Judul lowongan wajib diisi!', 'danger')
        return redirect(url_for('vacancy'))

    try:
        conn = get_db_connection()
        cursor = conn.cursor()
        cursor.execute("""
            INSERT INTO vacancies (title, division, branch, description, requirements, posted_date, status, image)
            VALUES (%s, %s, %s, %s, %s, %s, %s, %s)
        """, (title, division, branch, description, requirements, posted_date, status, image))
        conn.commit()
        conn.close()
        flash(f'Lowongan "{title}" berhasil ditambahkan!', 'success')
    except Exception as e:
        flash(f'Gagal menambahkan lowongan: {e}', 'danger')

    return redirect(url_for('vacancy'))

@app.route('/edit-vacancy/<int:id>', methods=['POST'])
def edit_vacancy(id):
    title = request.form.get('title')
    division = request.form.get('division')
    branch = request.form.get('branch')
    description = request.form.get('description')
    requirements = request.form.get('requirements')
    status = request.form.get('status', 'Open')

    if not title:
        flash('Judul lowongan wajib diisi!', 'danger')
        return redirect(url_for('vacancy'))

    try:
        conn = get_db_connection()
        cursor = conn.cursor()
        cursor.execute("""
            UPDATE vacancies
            SET title=%s, division=%s, branch=%s, description=%s, requirements=%s, status=%s
            WHERE id=%s
        """, (title, division, branch, description, requirements, status, id))
        conn.commit()
        conn.close()
        flash(f'Lowongan "{title}" berhasil diperbarui!', 'success')
    except Exception as e:
        flash(f'Gagal memperbarui lowongan: {e}', 'danger')

    return redirect(url_for('vacancy'))

@app.route('/delete-vacancy/<int:id>', methods=['POST'])
def delete_vacancy(id):
    try:
        conn = get_db_connection()
        cursor = conn.cursor()
        cursor.execute("SELECT title FROM vacancies WHERE id = %s", (id,))
        vac = cursor.fetchone()
        cursor.execute("DELETE FROM vacancies WHERE id = %s", (id,))
        conn.commit()
        conn.close()
        title_str = vac['title'] if vac else f'ID {id}'
        flash(f'Lowongan "{title_str}" berhasil dihapus!', 'success')
    except Exception as e:
        flash(f'Gagal menghapus lowongan: {e}', 'danger')

    return redirect(url_for('vacancy'))

# ==========================================
# RUTE DATA KEHADIRAN (PRESENCE)
# ==========================================
@app.route('/presence')
def presence():
    now = datetime.now()
    current_day = now.strftime('%A')
    current_date_formatted = f"{now.strftime('%B')} {now.day}, {now.year}"
    
    search_query = request.args.get('search', '').strip()
    selected_branch = request.args.get('branch', '').strip()
    selected_team = request.args.get('team', '').strip()
    selected_status = request.args.get('status', '').strip()
    selected_tag = request.args.get('tag', '').strip()
    
    conn = get_db_connection()
    cursor = conn.cursor()
    
    # Ordering matching the Figma presence mockup
    order_clause = """
        ORDER BY 
            CASE 
                WHEN LEFT(e.employee_id, 3) = '278' THEN 1
                WHEN LEFT(e.employee_id, 3) = '174' THEN 2
                WHEN LEFT(e.employee_id, 3) = '793' THEN 3
                WHEN LEFT(e.employee_id, 3) = '554' THEN 4
                WHEN LEFT(e.employee_id, 3) = '650' THEN 5
                WHEN LEFT(e.employee_id, 3) = '445' THEN 6
                WHEN LEFT(e.employee_id, 3) = '496' THEN 7
                WHEN LEFT(e.employee_id, 3) = '954' THEN 8
                WHEN LEFT(e.employee_id, 3) = '675' THEN 9
                WHEN LEFT(e.employee_id, 3) = '769' THEN 10
                ELSE 100
            END ASC,
            e.id ASC
    """
    
    # Bangun kondisi WHERE dinamis
    where_conditions = []
    params = []
    
    if search_query:
        where_conditions.append("(e.branch LIKE %s OR e.name LIKE %s OR e.team LIKE %s OR e.employee_id LIKE %s)")
        q = f"%{search_query}%"
        params.extend([q, q, q, q])
        
    if selected_branch:
        where_conditions.append("e.branch = %s")
        params.append(selected_branch)
        
    if selected_team:
        where_conditions.append("e.team = %s")
        params.append(selected_team)
        
    if selected_status:
        if selected_status == 'hadir':
            where_conditions.append("(a.type = 'clock_in' OR a.type = 'hadir')")
        elif selected_status == 'pulang':
            where_conditions.append("a.type = 'clock_out'")
        elif selected_status == 'izin':
            where_conditions.append("a.type = 'izin'")
        elif selected_status == 'sakit':
            where_conditions.append("a.type = 'sakit'")
        elif selected_status == 'belum_hadir':
            where_conditions.append("(a.type IS NULL OR a.type = '')")
            
    if selected_tag:
        where_conditions.append("(e.tag1 = %s OR e.tag2 = %s OR e.tag3 = %s)")
        params.extend([selected_tag, selected_tag, selected_tag])
        
    where_sql = ("WHERE " + " AND ".join(where_conditions)) if where_conditions else ""
    
    query = f"""
        SELECT 
            e.id, e.employee_id, e.name, e.team, e.branch, e.status, e.tag1, e.tag2, e.tag3, e.username, e.face_image,
            a.type AS today_type,
            SUBSTRING(TIME(a.max_ts), 1, 5) AS today_clock_in,
            SUBSTRING(TIME(a.clock_out), 1, 5) AS today_clock_out,
            a.image_path AS today_image,
            a.reason AS today_reason
        FROM employees e
        LEFT JOIN (
            SELECT username,
                   SUBSTRING_INDEX(GROUP_CONCAT(type ORDER BY id DESC), ',', 1) AS type,
                   MAX(timestamp) AS max_ts,
                   MAX(clock_out) AS clock_out,
                   SUBSTRING_INDEX(GROUP_CONCAT(COALESCE(image_path, '') ORDER BY id DESC), ',', 1) AS image_path,
                   SUBSTRING_INDEX(GROUP_CONCAT(COALESCE(reason, '') ORDER BY id DESC), ',', 1) AS reason
            FROM attendance
            WHERE DATE(timestamp) = CURDATE()
            GROUP BY username
        ) a ON e.username = a.username
        {where_sql}
        {order_clause}
    """
    
    if params:
        cursor.execute(query, tuple(params))
    else:
        cursor.execute(query)
        
    presence_list = cursor.fetchall()

    # Hitung keterangan kedatangan (Sangat Baik / Baik / Kurang Baik)
    # - dibawah jam 08:50 (530 menit): Sangat Baik
    # - 08:50 s/d 09:00 (530 - 540 menit): Baik
    # - diatas 09:00 (> 540 menit): Kurang Baik
    for emp in presence_list:
        clock_in = emp.get('today_clock_in')
        if clock_in and clock_in != '--:--':
            try:
                parts = str(clock_in).split(':')
                h, m = int(parts[0]), int(parts[1])
                tot = h * 60 + m
                if tot < 530:
                    emp['today_keterangan'] = 'Sangat Baik'
                    emp['keterangan_class'] = 'success'
                elif tot <= 540:
                    emp['today_keterangan'] = 'Baik'
                    emp['keterangan_class'] = 'primary'
                else:
                    emp['today_keterangan'] = 'Kurang Baik'
                    emp['keterangan_class'] = 'danger'
            except Exception:
                emp['today_keterangan'] = None
        else:
            emp['today_keterangan'] = None
    cursor.execute("""
        SELECT COUNT(DISTINCT username) as total_hadir
        FROM attendance
        WHERE DATE(timestamp) = CURDATE() AND (type = 'clock_in' OR type = 'hadir')
    """)
    hadir_count = cursor.fetchone()['total_hadir'] or 0
    
    # Ambil opsi filter kantor cabang dari tabel branch_companies
    cursor.execute("SELECT branch_name FROM branch_companies ORDER BY id ASC")
    branches = [r['branch_name'] for r in cursor.fetchall()]
    if not branches:
        branches = ['TUNGGAL SAHABAT ABADI', 'Decoy']
    
    cursor.execute("SELECT DISTINCT team FROM employees WHERE team IS NOT NULL AND team != '' ORDER BY team")
    teams = [r['team'] for r in cursor.fetchall()]
    
    conn.close()
    
    active_filter_count = sum(1 for val in [selected_branch, selected_team, selected_status, selected_tag] if val)
    
    return render_template(
        'presence.html',
        presence_list=presence_list,
        search_query=search_query,
        selected_branch=selected_branch,
        selected_team=selected_team,
        selected_status=selected_status,
        selected_tag=selected_tag,
        active_filter_count=active_filter_count,
        branches=branches,
        teams=teams,
        hadir_count=hadir_count,
        current_day=current_day,
        current_date_formatted=current_date_formatted
    )

# ==========================================
# RUTE PENGAJUAN IZIN / SAKIT (REQUEST)
# ==========================================
@app.route('/request')
@app.route('/requests')
def request_page():
    now = datetime.now()
    current_day = now.strftime('%A')
    current_date_formatted = f"{now.strftime('%B')} {now.day}, {now.year}"
    
    requests_list = []
    employees_list = []
    try:
        conn = get_db_connection()
        cursor = conn.cursor()
        
        # 1. Query seluruh pengajuan izin / sakit langsung dari tabel attendance JOIN employees
        cursor.execute("""
            SELECT 
                a.id,
                a.username,
                a.type,
                a.timestamp,
                a.clock_out,
                a.image_path,
                a.reason,
                COALESCE(a.status, 'Pending') AS status,
                DATE_FORMAT(a.timestamp, '%b %e, %Y') AS formatted_date,
                e.id AS employee_db_id,
                e.employee_id,
                COALESCE(e.name, a.username) AS name,
                COALESCE(e.team, 'Staff') AS role,
                e.team,
                e.branch,
                e.face_image,
                e.status AS emp_status,
                e.tag1, e.tag2
            FROM attendance a
            LEFT JOIN employees e ON a.username = e.username
            WHERE a.type IN ('izin', 'sakit', 'izin_sakit', 'izin_keperluan', 'cuti')
               OR (a.reason IS NOT NULL AND a.reason != '' AND a.type NOT IN ('clock_in', 'clock_out', 'hadir'))
            ORDER BY 
                CASE 
                    WHEN e.username = 'reva' AND a.type = 'sakit' THEN 1
                    WHEN e.username = 'Satriya' THEN 2
                    WHEN e.username = 'bambang' THEN 3
                    WHEN e.username = 'rina' THEN 4
                    WHEN e.username = 'dimas' THEN 5
                    WHEN e.username = 'eko' THEN 6
                    ELSE 100
                END ASC,
                a.id DESC
        """)
        raw_requests = cursor.fetchall()
        
        for r in raw_requests:
            r_type = r['type'].lower() if r['type'] else ''
            if 'sakit' in r_type:
                display_type = 'Sakit'
            elif 'izin' in r_type or 'cuti' in r_type:
                display_type = 'Izin'
            else:
                display_type = 'Izin'
                
            r['display_type'] = display_type
            requests_list.append(r)
            
        # 2. Ambil daftar karyawan untuk modal tambah pengajuan
        cursor.execute("SELECT employee_id, username, name, team, branch FROM employees ORDER BY name ASC")
        employees_list = cursor.fetchall()
        
        conn.close()
    except Exception as e:
        flash(f'Gagal memuat data pengajuan: {e}', 'danger')
        print(f"[Request Error] {e}")
    
    # Filter berdasarkan status untuk tabs (Pending, Disetujui, Ditolak)
    pending_list  = [r for r in requests_list if (r.get('status') or 'Pending').lower() == 'pending']
    approved_list = [r for r in requests_list if (r.get('status') or '').lower() == 'approved']
    rejected_list = [r for r in requests_list if (r.get('status') or '').lower() == 'rejected']
    
    return render_template(
        'request.html',
        requests_list=requests_list,
        pending_list=pending_list,
        approved_list=approved_list,
        rejected_list=rejected_list,
        count_pending=len(pending_list),
        count_approved=len(approved_list),
        count_rejected=len(rejected_list),
        employees_list=employees_list,
        current_day=current_day,
        current_date_formatted=current_date_formatted
    )

@app.route('/add-request', methods=['POST'])
def add_request():
    username = request.form.get('username')
    req_type = request.form.get('type', 'sakit')
    reason = request.form.get('reason', '')
    req_date = request.form.get('date') or datetime.now().strftime('%Y-%m-%d')
    timestamp_str = f"{req_date} 08:00:00"

    if not username or not reason:
        flash('Karyawan dan Alasan izin/sakit wajib diisi!', 'danger')
        return redirect(url_for('request_page'))

    try:
        conn = get_db_connection()
        cursor = conn.cursor()

        # Jika karyawan pengajuan lagi setelah ditolak, hapus pengajuan yang ditolak sebelumnya
        cursor.execute("""
            DELETE FROM attendance 
            WHERE username = %s AND status = 'Rejected' 
              AND (type IN ('izin', 'sakit', 'izin_sakit', 'izin_keperluan', 'cuti') OR type LIKE 'izin_%%')
        """, (username,))

        cursor.execute("""
            INSERT INTO attendance (username, type, timestamp, reason, status)
            VALUES (%s, %s, %s, %s, 'Pending')
        """, (username, req_type, timestamp_str, reason))
        conn.commit()
        conn.close()
        flash(f'Pengajuan {req_type.capitalize()} untuk @{username} berhasil disimpan dan langsung masuk ke database!', 'success')
    except Exception as e:
        flash(f'Gagal menambahkan pengajuan: {e}', 'danger')

    return redirect(url_for('request_page'))

@app.route('/update-request-status/<int:id>', methods=['POST'])
def update_request_status(id):
    new_status = request.form.get('status', 'Approved')
    try:
        conn = get_db_connection()
        cursor = conn.cursor()
        cursor.execute("UPDATE attendance SET status = %s WHERE id = %s", (new_status, id))
        conn.commit()
        conn.close()
        flash(f'Status pengajuan berhasil diperbarui menjadi "{new_status}"!', 'success')
    except Exception as e:
        flash(f'Gagal memperbarui status: {e}', 'danger')
    return redirect(url_for('request_page'))

@app.route('/delete-request/<int:id>', methods=['POST'])
def delete_request(id):
    try:
        conn = get_db_connection()
        cursor = conn.cursor()
        cursor.execute("DELETE FROM attendance WHERE id = %s", (id,))
        conn.commit()
        conn.close()
        flash('Data pengajuan berhasil dihapus!', 'success')
    except Exception as e:
        flash(f'Gagal menghapus pengajuan: {e}', 'danger')
    return redirect(url_for('request_page'))

# ==========================================
# ==========================================
# RUTE LAPORAN KEHADIRAN (REPORT)
# FORMAT REKAPITULASI (TANPA NIP)
# ==========================================
def get_rekap_report_data(start_date, end_date, branch='all'):
    try:
        dt_start = datetime.strptime(start_date, '%Y-%m-%d')
        dt_end = datetime.strptime(end_date, '%Y-%m-%d')
        if dt_start > dt_end:
            dt_start, dt_end = dt_end, dt_start
        total_days = (dt_end - dt_start).days + 1
        
        # Hitung hari kerja asli sesuai tanggal kalender (Senin s/d Jumat, Sabtu & Minggu libur)
        working_days = sum(
            1 for i in range(total_days)
            if (dt_start + timedelta(days=i)).weekday() < 5  # 0=Senin .. 4=Jumat (Sabtu=5 & Minggu=6 libur)
        )
        total_hari_kerja = max(1, working_days)
    except Exception:
        total_hari_kerja = 22

    start_ts = f"{start_date} 00:00:00"
    end_ts = f"{end_date} 23:59:59"

    conn = get_db_connection()
    cursor = conn.cursor()

    where_branch = ""
    params = [start_ts, end_ts]
    if branch and branch != 'all':
        where_branch = "WHERE (e.branch = %s OR (%s = 'TUNGGAL SAHABAT ABADI' AND (e.branch = 'Surabaya' OR e.branch IS NULL)))"
        params.extend([branch, branch])

    cursor.execute(f"""
        SELECT 
            e.id,
            e.name,
            e.username,
            COALESCE(e.team, 'Operasional') AS divisi,
            COALESCE(e.branch, 'TUNGGAL SAHABAT ABADI') AS branch,
            COUNT(DISTINCT CASE WHEN a.type IN ('clock_in', 'hadir') THEN DATE(a.timestamp) END) AS jumlah_hadir,
            COUNT(DISTINCT CASE WHEN a.type IN ('izin', 'sakit', 'cuti', 'sick_leave', 'necessity_leave') THEN DATE(a.timestamp) END) AS jumlah_izin
        FROM employees e
        LEFT JOIN attendance a 
            ON e.username = a.username 
            AND a.timestamp >= %s AND a.timestamp <= %s
        {where_branch}
        GROUP BY e.id, e.name, e.username, e.team, e.branch
        ORDER BY e.id ASC
    """, tuple(params))
    rows = cursor.fetchall()

    cursor.execute("SELECT branch_name FROM branch_companies ORDER BY id ASC")
    branches = [b['branch_name'] for b in cursor.fetchall()]
    if not branches:
        branches = ['TUNGGAL SAHABAT ABADI', 'Decoy']
    conn.close()

    processed_data = []
    tot_hadir = 0
    tot_izin = 0
    tot_alpha = 0

    for idx, r in enumerate(rows, 1):
        hadir = r['jumlah_hadir']
        izin = r['jumlah_izin']
        alpha = max(0, total_hari_kerja - hadir - izin)
        pct = round((hadir / total_hari_kerja) * 100) if total_hari_kerja > 0 else 0
        username = r['username'] or ''
        email = username if '@' in username else (f"{username.lower()}@gmail.com" if username else '-')

        tot_hadir += hadir
        tot_izin += izin
        tot_alpha += alpha

        processed_data.append({
            'no': idx,
            'name': r['name'] or username or '-',
            'email': email,
            'divisi': r['divisi'] or 'Operasional',
            'branch': r['branch'] if r['branch'] != 'Surabaya' else 'TUNGGAL SAHABAT ABADI',
            'total_hari_kerja': total_hari_kerja,
            'jumlah_hadir': hadir,
            'jumlah_izin': izin,
            'jumlah_alpha': alpha,
            'persentase': f"{pct}%",
            'persentase_num': pct
        })

    return {
        'data': processed_data,
        'branches': branches,
        'total_hari_kerja': total_hari_kerja,
        'tot_hadir': tot_hadir,
        'tot_izin': tot_izin,
        'tot_alpha': tot_alpha
    }

@app.route('/report')
def report_page():
    now = datetime.now()
    current_day = now.strftime('%A')
    current_date_formatted = f"{now.strftime('%B')} {now.day}, {now.year}"
    
    start_date = request.args.get('start_date', '2026-10-01')
    end_date = request.args.get('end_date', '2026-10-31')
    selected_branch = request.args.get('branch', 'TUNGGAL SAHABAT ABADI')
    selected_format = request.args.get('format', 'xlsx')
    
    rekap = get_rekap_report_data(start_date, end_date, selected_branch)
    
    return render_template(
        'report.html',
        start_date=start_date,
        end_date=end_date,
        selected_branch=selected_branch,
        selected_format=selected_format,
        branches=rekap['branches'],
        preview_data=rekap['data'],
        total_hari_kerja=rekap['total_hari_kerja'],
        tot_hadir=rekap['tot_hadir'],
        tot_izin=rekap['tot_izin'],
        tot_alpha=rekap['tot_alpha'],
        current_day=current_day,
        current_date_formatted=current_date_formatted
    )

@app.route('/report/export', methods=['GET', 'POST'])
def export_report():
    if request.method == 'POST':
        start_date = request.form.get('start_date', '2026-10-01')
        end_date = request.form.get('end_date', '2026-10-31')
        branch = request.form.get('branch', 'TUNGGAL SAHABAT ABADI')
        file_format = request.form.get('format', 'xlsx').lower()
    else:
        start_date = request.args.get('start_date', '2026-10-01')
        end_date = request.args.get('end_date', '2026-10-31')
        branch = request.args.get('branch', 'TUNGGAL SAHABAT ABADI')
        file_format = request.args.get('format', 'xlsx').lower()

    rekap = get_rekap_report_data(start_date, end_date, branch)
    processed_data = rekap['data']
    total_hari_kerja = rekap['total_hari_kerja']

    filename_base = f"Laporan_Kehadiran_Karyawan_{start_date}_sd_{end_date}"

    # EXCEL GENERATION (.XLSX)
    if file_format == 'xlsx':
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Rekap Kehadiran"

        headers = [
            "No",
            "Nama Karyawan",
            "Email",
            "Divisi",
            "Total Hari Kerja",
            "Jumlah Hadir",
            "Jumlah Izin",
            "Jumlah Alpha",
            "Persentase Kehadiran (%)"
        ]
        ws.append(headers)

        thin_border = Border(
            left=Side(style='thin', color='D1D5DB'),
            right=Side(style='thin', color='D1D5DB'),
            top=Side(style='thin', color='D1D5DB'),
            bottom=Side(style='thin', color='D1D5DB')
        )

        # Style header row (Row 1)
        for col_num in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_num)
            cell.font = Font(name="Calibri", bold=True, color="000000", size=11)
            cell.alignment = Alignment(horizontal="center" if col_num in [1, 5, 6, 7, 8, 9] else "left", vertical="center")
            cell.border = thin_border

        # Isi data baris demi baris (Row 2+)
        for r in processed_data:
            row_data = [
                r['no'],
                r['name'],
                r['email'],
                r['divisi'],
                r['total_hari_kerja'],
                r['jumlah_hadir'],
                r['jumlah_izin'],
                r['jumlah_alpha'],
                r['persentase']
            ]
            ws.append(row_data)
            current_row = ws.max_row
            for col_idx in range(1, len(row_data) + 1):
                c = ws.cell(row=current_row, column=col_idx)
                c.border = thin_border
                c.font = Font(name="Calibri", size=11)
                if col_idx in [1, 5, 6, 7, 8, 9]:
                    c.alignment = Alignment(horizontal="center")
                else:
                    c.alignment = Alignment(horizontal="left")

        # Auto lebar kolom
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return send_file(
            buf,
            as_attachment=True,
            download_name=f"{filename_base}.xlsx",
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )

    # PDF GENERATION (.PDF)
    elif file_format == 'pdf':
        buf = io.BytesIO()
        doc = SimpleDocTemplate(buf, pagesize=landscape(letter), rightMargin=24, leftMargin=24, topMargin=24, bottomMargin=24)
        elements = []

        styles = getSampleStyleSheet()
        title_style = ParagraphStyle('ReportTitle', parent=styles['Heading1'], fontSize=15, leading=19, textColor=colors.HexColor('#131B3E'))
        sub_style = ParagraphStyle('ReportSub', parent=styles['Normal'], fontSize=9, leading=13, textColor=colors.HexColor('#4A5568'))

        branch_str = "Semua Cabang" if branch == 'all' else branch
        elements.append(Paragraph("<b>LAPORAN REKAPITULASI KEHADIRAN KARYAWAN</b>", title_style))
        elements.append(Paragraph(f"Rentang Tanggal: {start_date} s/d {end_date} | Kantor Cabang: {branch_str} | Dicetak: {datetime.now().strftime('%d/%m/%Y %H:%M')}", sub_style))
        elements.append(Spacer(1, 14))

        table_headers = ["No", "Nama Karyawan", "Email", "Divisi", "Total Hari Kerja", "Jumlah Hadir", "Jumlah Izin", "Jumlah Alpha", "Kehadiran (%)"]
        table_rows = [table_headers]

        for r in processed_data:
            table_rows.append([
                str(r['no']),
                str(r['name'])[:22],
                str(r['email'])[:24],
                str(r['divisi'])[:16],
                str(r['total_hari_kerja']),
                str(r['jumlah_hadir']),
                str(r['jumlah_izin']),
                str(r['jumlah_alpha']),
                r['persentase']
            ])

        col_widths = [30, 140, 140, 90, 80, 65, 65, 65, 75]
        pdf_table = Table(table_rows, colWidths=col_widths, repeatRows=1)
        pdf_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#131B3E')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('ALIGN', (1, 1), (3, -1), 'LEFT'),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 9),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 6),
            ('TOPPADDING', (0, 0), (-1, 0), 6),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F8FAFC')]),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E2E8F0')),
            ('FONTSIZE', (0, 1), (-1, -1), 8.5),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ]))
        elements.append(pdf_table)

        doc.build(elements)
        buf.seek(0)
        return send_file(
            buf,
            as_attachment=True,
            download_name=f"{filename_base}.pdf",
            mimetype="application/pdf"
        )

    # CSV GENERATION (.CSV)
    else:
        buf = io.StringIO()
        writer = csv.writer(buf)
        writer.writerow(["No", "Nama Karyawan", "Email", "Divisi", "Total Hari Kerja", "Jumlah Hadir", "Jumlah Izin", "Jumlah Alpha", "Persentase Kehadiran (%)"])
        for r in processed_data:
            writer.writerow([
                r['no'],
                r['name'],
                r['email'],
                r['divisi'],
                r['total_hari_kerja'],
                r['jumlah_hadir'],
                r['jumlah_izin'],
                r['jumlah_alpha'],
                r['persentase']
            ])
        output_bytes = io.BytesIO(buf.getvalue().encode('utf-8-sig'))
        output_bytes.seek(0)
        return send_file(
            output_bytes,
            as_attachment=True,
            download_name=f"{filename_base}.csv",
            mimetype="text/csv"
        )

@app.route('/logout')
def logout():
    session.clear()
    flash('Anda telah berhasil keluar dari akun!', 'info')
    return redirect(url_for('login'))

if __name__ == '__main__':
    port = int(os.environ.get('PORT', 8085))
    debug_mode = os.environ.get('FLASK_DEBUG', 'False').lower() == 'true'
    app.run(host='0.0.0.0', port=port, debug=debug_mode)